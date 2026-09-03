from pathlib import Path

import re
import sys
import traceback
from pypdf import PdfReader
from docx import Document
from embedder import Embedder
from database import Database
import requests
from bs4 import BeautifulSoup


KNOWN_FILIERES = {
    "SDBDIA",
    "SITCN",
    "MGSI",
    "IL",
    "ENSIASD"
}


def detect_filiere(name):
    if not name:
        return None
    name_upper = name.upper()
    for filiere in KNOWN_FILIERES:
        if re.search(rf"\b{filiere}\b", name_upper) or filiere in name_upper:
            return filiere
    return None


def detect_semester(text):
    if not text:
        return None

    match = re.search(
        r"semestre\s*(?:de\s*programmation\s*(?:du\s*module)?)?\s*[:\s]*([1-8])\b",
        text,
        flags=re.IGNORECASE
    )
    if match:
        return f"S{match.group(1)}"

    match = re.search(
        r"\bS([1-8])\b",
        text,
        flags=re.IGNORECASE
    )
    if match:
        return f"S{match.group(1)}"

    return None


def detect_module_code(text):
    if not text:
        return None

    match = re.search(
        r"\b(M\d{3}(?:[_\-.]\d+)?)\b",
        text,
        flags=re.IGNORECASE
    )

    if match:
        code = match.group(1).upper()
        code = re.sub(r"[-.]", "_", code)
        return code

    return None


def detect_module_title(text):
    if not text:
        return None

    match = re.search(
        r"Intitul[ée]\s*(?:du\s*module|Module)?\s*[:\s]*([^\n\r]+)",
        text,
        flags=re.IGNORECASE
    )

    if match:
        title = match.group(1).strip()
        title = re.sub(r"\s+", " ", title)
        title_lower = title.lower()

        if (
            len(title) > 3
            and not title_lower.startswith("neant")
            and not title_lower.startswith("néant")
            and not title_lower.startswith("disciplinaire")
            and not title_lower.startswith("transversal")
        ):
            return title

    return None


def detect_section_heading(line):
    if not line:
        return None

    clean = line.strip()

    heading_patterns = [
        r"^(?:\d+[\.\)]\s*)?DESCRIPTION\s+(?:DU\s+CONTENU|SUCCINCTE|DETAILLEE)\b.*",
        r"^(?:\d+[\.\)]\s*)?ORGANISATION\s+MODULAIRE\b.*",
        r"^(?:\d+[\.\)]\s*)?CALENDRIER\s+DES\s+EXAMENS\b.*",
        r"^(?:\d+[\.\)]\s*)?(?:COORDONNATEUR\s+ET\s+)?EQUIPE\s+PEDAGOGIQUE\b.*",
        r"^(?:\d+[\.\)]\s*)?VOLUME\s+HORAIRE\b.*",
        r"^(?:\d+[\.\)]\s*)?MODALITES\s+D['’]EVALUATION\b.*",
        r"^(?:\d+[\.\)]\s*)?EVALUATION\s+DU\s+MODULE\b.*",
        r"^(?:\d+[\.\)]\s*)?VALIDATION\s+DU\s+MODULE\b.*",
        r"^(?:\d+[\.\)]\s*)?PR[EÉ]REQUIS\s+P[EÉ]DAGOGIQUES\b.*",
        r"^(?:\d+[\.\)]\s*)?OBJECTIFS?\s+DU\s+MODULE\b.*",
        r"^(?:\d+[\.\)]\s*)?CONDITIONS\s+(?:ET\s+MODALIT[EÉ]S\s+)?D['’]ACC[EÈ]S\b.*",
        r"^(?:\d+[\.\)]\s*)?D[EÉ]BOUCH[EÉ]S\b.*",
        r"^(?:\d+[\.\)]\s*)?R[EÈ]GLEMENT\s+INT[EÉ]RIEUR\b.*",
        r"^(?:\d+[\.\)]\s*)?SOMMAIRE\s+DES\s+DESCRIPTIFS\b.*",
        r"^DESCRIPTIF\s+DU\s+MODULE\b.*",
    ]

    for pattern in heading_patterns:
        if re.search(pattern, clean, flags=re.IGNORECASE):
            return clean

    return None


def extract_page_module_context(page_text):
    context = {}

    m_code = re.search(
        r"(?:N[°o]\s*d['’]ordre\s*du\s*module|Code\s*du\s*module)[\s\S]{0,50}?\b(M\d{3}(?:[_\-.]\d+)?)\b",
        page_text,
        flags=re.IGNORECASE
    )

    if m_code:
        code = m_code.group(1).upper()
        context["module_code"] = re.sub(r"[-.]", "_", code)
    else:
        m_code_fallback = re.search(
            r"\b(M\d{3}(?:[_\-.]\d+)?)\b",
            page_text,
            flags=re.IGNORECASE
        )

        if m_code_fallback and "DESCRIPTIF DU MODULE" in page_text.upper():
            code = m_code_fallback.group(1).upper()
            context["module_code"] = re.sub(r"[-.]", "_", code)

    m_title = re.search(
        r"Intitul[ée]\s*(?:du\s*module|Module)?\s*[:\s]*([^\n\r]+)",
        page_text,
        flags=re.IGNORECASE
    )

    if m_title:
        title = m_title.group(1).strip()
        title = re.sub(r"\s+", " ", title)
        title_lower = title.lower()

        if (
            len(title) > 3
            and not title_lower.startswith("neant")
            and not title_lower.startswith("néant")
            and not title_lower.startswith("disciplinaire")
            and not title_lower.startswith("transversal")
        ):
            context["module_title"] = title

    m_sem = re.search(
        r"Semestre\s*(?:de\s*programmation\s*(?:du\s*module)?)?[\s\S]{0,40}?\b([1-8])\b",
        page_text,
        flags=re.IGNORECASE
    )

    if m_sem:
        context["semester"] = f"S{m_sem.group(1)}"

    return context


def build_context_header(
    filiere=None,
    semester=None,
    module_code=None,
    module_title=None,
    heading=None
):
    parts = []

    if filiere:
        parts.append(f"Filière: {filiere}")

    if semester:
        parts.append(f"Semestre: {semester}")

    if module_code:
        if module_title:
            parts.append(f"Module: {module_code} - {module_title}")
        else:
            parts.append(f"Module: {module_code}")
    elif module_title:
        parts.append(f"Module: {module_title}")

    if heading:
        parts.append(f"Section: {heading}")

    if not parts:
        return ""

    return " | ".join(parts)


def extract_pdf(file_path):
    reader = PdfReader(file_path)

    if len(reader.pages) == 0:
        print(f"  WARNING: PDF has 0 pages: {file_path}")
        return []

    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text(extraction_mode="layout")
        except TypeError:
            try:
                text = page.extract_text()
            except Exception as e:
                print(f"  WARNING: Could not extract page {page_number}: {e}")
                continue
        except Exception as e:
            print(f"  WARNING: Error extracting page {page_number}: {e}")
            continue

        if not text:
            continue

        text = clean_extracted_text(text)

        if text.strip():
            pages.append({
                "page": page_number,
                "text": text
            })

    return pages


def clean_extracted_text(text):
    lines = []

    for line in text.splitlines():
        line = line.replace("\xa0", " ")
        line = re.sub(r" {4,}", "   ", line)
        line = line.strip()

        if not line:
            if lines and lines[-1] != "":
                lines.append("")
            continue

        lines.append(line)

    while lines and lines[-1] == "":
        lines.pop()

    return "\n".join(lines)


def extract_docx(file_path):
    document = Document(file_path)
    content_parts = []

    for element in document.element.body:
        tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag

        if tag == "p":
            text = element.text
            if text and text.strip():
                content_parts.append(text.strip())

        elif tag == "tbl":
            table_lines = _extract_docx_table_element(element)
            if table_lines:
                content_parts.extend(table_lines)

    if not content_parts:
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                content_parts.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    content_parts.append(" | ".join(cells))

    if not content_parts:
        return []

    return [{
        "page": None,
        "text": "\n".join(content_parts)
    }]


def _extract_docx_table_element(tbl_element):
    lines = []

    try:
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

        for tr in tbl_element.findall(f".//{ns}tr"):
            cells = []

            for tc in tr.findall(f".//{ns}tc"):
                cell_text = ""

                for p in tc.findall(f".//{ns}p"):
                    if p.text:
                        cell_text += p.text

                    for r in p.findall(f".//{ns}r"):
                        for t in r.findall(f".//{ns}t"):
                            if t.text:
                                cell_text += t.text

                cell_text = cell_text.strip()

                if cell_text:
                    cells.append(cell_text)

            if cells:
                lines.append(" | ".join(cells))

    except Exception:
        pass

    return lines


def extract_web(url):
    response = requests.get(
        url,
        timeout=30
    )

    response.raise_for_status()

    soup = BeautifulSoup(
        response.text,
        "html.parser"
    )

    text = soup.get_text(
        "\n",
        strip=True
    )

    return [
        {
            "page": None,
            "text": text
        }
    ]


def split_text(
    text,
    chunk_size=1200,
    overlap=250
):
    if not text or not text.strip():
        return []

    text = text.strip()

    if len(text) <= chunk_size:
        return [text]

    lines = text.splitlines()
    chunks = []
    current_lines = []
    current_length = 0

    for line in lines:
        line = line.strip()

        if not line:
            continue

        line_length = len(line)

        if current_lines and current_length + line_length + 1 > chunk_size:
            chunk = "\n".join(current_lines).strip()

            if chunk:
                chunks.append(chunk)

            overlap_lines = []
            overlap_length = 0

            for previous_line in reversed(current_lines):
                if overlap_length + len(previous_line) + 1 > overlap:
                    break

                overlap_lines.insert(0, previous_line)
                overlap_length += len(previous_line) + 1

            current_lines = overlap_lines
            current_length = overlap_length

        current_lines.append(line)
        current_length += line_length + 1

    if current_lines:
        chunk = "\n".join(current_lines).strip()

        if chunk:
            chunks.append(chunk)

    return chunks


def build_document_chunks(
    pages,
    document_name,
    chunk_size=1200,
    overlap=250
):
    doc_filiere = detect_filiere(document_name)
    doc_semester = detect_semester(document_name)

    all_lines = []

    current_filiere = doc_filiere
    current_semester = doc_semester
    current_module_code = None
    current_module_title = None
    current_heading = None

    for page in pages:
        page_num = page["page"]
        page_text = page["text"]

        if not page_text.strip():
            continue

        page_ctx = extract_page_module_context(page_text)

        if "module_code" in page_ctx:
            current_module_code = page_ctx["module_code"]

        if "module_title" in page_ctx:
            current_module_title = page_ctx["module_title"]

        if "semester" in page_ctx and not doc_semester:
            current_semester = page_ctx["semester"]

        for line in page_text.splitlines():
            clean = line.strip()

            if not clean:
                continue

            if "DESCRIPTIF DU MODULE" in clean.upper():
                current_heading = None

            detected_sem = detect_semester(clean)

            if detected_sem and not doc_semester:
                if (
                    "semestre" in clean.lower()
                    or re.match(r"^S[1-8]$", clean, re.IGNORECASE)
                ):
                    current_semester = detected_sem

            detected_code = detect_module_code(clean)

            if detected_code:
                if (
                    "module" in clean.lower()
                    or re.match(
                        r"^M\d{3}(?:[_\-.]\d+)?\b",
                        clean,
                        re.IGNORECASE
                    )
                ):
                    current_module_code = detected_code

            detected_title = detect_module_title(clean)

            if detected_title:
                current_module_title = detected_title

            detected_hd = detect_section_heading(clean)

            if detected_hd:
                current_heading = detected_hd

            all_lines.append({
                "text": clean,
                "page": page_num,
                "filiere": current_filiere,
                "semester": current_semester,
                "module_code": current_module_code,
                "module_title": current_module_title,
                "heading": current_heading
            })

    if not all_lines:
        return []

    chunks = []
    current_items = []
    current_length = 0
    current_pages = set()

    for item in all_lines:
        line_len = len(item["text"])

        if current_items and current_length + line_len + 1 > chunk_size:
            raw_content = "\n".join(
                i["text"] for i in current_items
            ).strip()

            if raw_content:
                c_filiere = next(
                    (i["filiere"] for i in current_items if i["filiere"]),
                    doc_filiere
                )

                c_sem = next(
                    (i["semester"] for i in current_items if i["semester"]),
                    doc_semester
                )

                c_code = next(
                    (i["module_code"] for i in current_items if i["module_code"]),
                    None
                )

                c_title = next(
                    (i["module_title"] for i in current_items if i["module_title"]),
                    None
                )

                c_heading = next(
                    (i["heading"] for i in current_items if i["heading"]),
                    None
                )

                header = build_context_header(
                    filiere=c_filiere,
                    semester=c_sem,
                    module_code=c_code,
                    module_title=c_title,
                    heading=c_heading
                )

                if header:
                    stored_content = f"[{header}]\n\n{raw_content}"
                else:
                    stored_content = raw_content

                pages_sorted = sorted(
                    p for p in current_pages if p is not None
                )

                chunks.append({
                    "content": stored_content,
                    "pages": pages_sorted,
                    "filiere": c_filiere,
                    "semester": c_sem,
                    "module_code": c_code,
                    "heading": c_heading
                })

            overlap_items = []
            overlap_len = 0

            for prev in reversed(current_items):
                p_len = len(prev["text"]) + 1

                if overlap_len + p_len > overlap:
                    break

                overlap_items.insert(0, prev)
                overlap_len += p_len

            current_items = overlap_items
            current_length = sum(
                len(i["text"]) + 1 for i in current_items
            )
            current_pages = {
                i["page"] for i in current_items
            }

        current_items.append(item)
        current_length += line_len + 1
        current_pages.add(item["page"])

    if current_items:
        raw_content = "\n".join(
            i["text"] for i in current_items
        ).strip()

        if raw_content:
            c_filiere = next(
                (i["filiere"] for i in current_items if i["filiere"]),
                doc_filiere
            )

            c_sem = next(
                (i["semester"] for i in current_items if i["semester"]),
                doc_semester
            )

            c_code = next(
                (i["module_code"] for i in current_items if i["module_code"]),
                None
            )

            c_title = next(
                (i["module_title"] for i in current_items if i["module_title"]),
                None
            )

            c_heading = next(
                (i["heading"] for i in current_items if i["heading"]),
                None
            )

            header = build_context_header(
                filiere=c_filiere,
                semester=c_sem,
                module_code=c_code,
                module_title=c_title,
                heading=c_heading
            )

            if header:
                stored_content = f"[{header}]\n\n{raw_content}"
            else:
                stored_content = raw_content

            pages_sorted = sorted(
                p for p in current_pages if p is not None
            )

            chunks.append({
                "content": stored_content,
                "pages": pages_sorted,
                "filiere": c_filiere,
                "semester": c_sem,
                "module_code": c_code,
                "heading": c_heading
            })

    return chunks


def build_embedding_text(
    document_name,
    pages,
    chunk
):
    lines = [f"Document: {document_name}"]

    if pages:
        if len(pages) == 1:
            lines.append(f"Page {pages[0]}")
        else:
            lines.append(
                "Pages " + ", ".join(str(p) for p in pages)
            )

    meta_parts = []

    filiere = chunk.get("filiere")
    semester = chunk.get("semester")
    module_code = chunk.get("module_code")
    heading = chunk.get("heading")

    if filiere:
        meta_parts.append(f"Filière: {filiere}")

    if semester:
        meta_parts.append(f"Semestre: {semester}")

    if module_code:
        meta_parts.append(f"Module: {module_code}")

    if heading:
        meta_parts.append(f"Section: {heading}")

    if meta_parts:
        lines.append(" | ".join(meta_parts))

    return "\n".join(lines) + "\n\n" + chunk["content"]


def ingest_file(file_path):
    file_path = Path(file_path)

    if not file_path.exists():
        print(f"  ERROR: File does not exist: {file_path}")
        return 0

    if file_path.stat().st_size == 0:
        print(f"  ERROR: File is empty (0 bytes): {file_path}")
        return 0

    extension = file_path.suffix.lower()

    embedder = Embedder()
    database = Database()

    try:
        if extension == ".pdf":
            pages = extract_pdf(file_path)
        elif extension == ".docx":
            pages = extract_docx(file_path)
        else:
            print(f"  Unsupported file type: {extension}")
            return 0

    except Exception as e:
        print(
            f"  ERROR: Extraction failed for {file_path.name}: {e}"
        )
        traceback.print_exc()
        return 0

    if not pages:
        print(
            f"  WARNING: No content extracted from {file_path.name}"
        )
        return 0

    document_name = file_path.name
    document_id = file_path.stem

    database.delete_document_chunks(document_id)

    chunks = build_document_chunks(
        pages=pages,
        document_name=document_name,
        chunk_size=1200,
        overlap=250
    )

    if not chunks:
        print(
            f"  WARNING: 0 chunks generated for {document_name}"
        )
        return 0

    inserted = 0

    for chunk_index, chunk in enumerate(chunks):
        pages_used = chunk["pages"]
        page_number = pages_used[0] if pages_used else None

        metadata = {
            "page": page_number,
            "pages": pages_used,
            "document": document_name,
            "filiere": chunk["filiere"],
            "semester": chunk["semester"],
            "module_code": chunk["module_code"],
            "heading": chunk["heading"]
        }

        embedding_text = build_embedding_text(
            document_name=document_name,
            pages=pages_used,
            chunk=chunk
        )

        try:
            chunk_embedding = embedder.get_embedding(
                embedding_text
            )

            database.insert_chunk(
                document_id=document_id,
                document_name=document_name,
                chunk_index=chunk_index,
                content=chunk["content"],
                embedding=chunk_embedding,
                source=document_name,
                metadata=metadata
            )

            inserted += 1

        except Exception as e:
            print(
                f"  ERROR: Failed to embed/insert chunk "
                f"{chunk_index} of {document_name}: {e}"
            )

    print(
        f"  Extracted {len(pages)} pages -> "
        f"Inserted {inserted}/{len(chunks)} chunks for {document_name}"
    )

    return inserted


def ingest_web(url):
    embedder = Embedder()
    database = Database()

    document_id = url
    document_name = url
    source = url

    try:
        pages = extract_web(url)
    except Exception as e:
        print(f"ERROR: Failed to fetch {url}: {e}")
        return

    database.delete_document_chunks(document_id)

    global_chunk_index = 0

    for page in pages:
        page_text = page["text"]

        if not page_text.strip():
            continue

        chunks = split_text(
            page_text,
            chunk_size=1200,
            overlap=250
        )

        for chunk in chunks:
            chunk_obj = {
                "content": chunk,
                "filiere": None,
                "semester": None,
                "module_code": None,
                "heading": None
            }

            embedding_text = build_embedding_text(
                document_name=document_name,
                pages=[],
                chunk=chunk_obj
            )

            chunk_embedding = embedder.get_embedding(
                embedding_text
            )

            metadata = {
                "page": None,
                "pages": [],
                "document": document_name,
                "filiere": None,
                "semester": None,
                "module_code": None,
                "heading": None
            }

            database.insert_chunk(
                document_id=document_id,
                document_name=document_name,
                chunk_index=global_chunk_index,
                content=chunk,
                embedding=chunk_embedding,
                source=source,
                metadata=metadata
            )

            global_chunk_index += 1

    print(f"Ingested {global_chunk_index} chunks from {url}")


def ingest_folder(folder_path):
    folder = Path(folder_path)

    if not folder.exists():
        print(f"ERROR: Folder does not exist: {folder_path}")
        return

    supported_extensions = {".pdf", ".docx"}

    files = sorted([
        f for f in folder.iterdir()
        if f.is_file() and f.suffix.lower() in supported_extensions
    ])

    if not files:
        print(
            f"WARNING: No supported documents found in {folder_path}"
        )
        return

    print(f"\n{'=' * 60}")
    print("INGESTION START")
    print(f"{'=' * 60}")
    print(f"Folder: {folder.resolve()}")
    print(f"Documents found: {len(files)}")

    for f in files:
        print(
            f"  - {f.name} ({f.stat().st_size / 1024:.0f} KB)"
        )

    print(f"{'=' * 60}\n")

    total_inserted = 0
    total_errors = 0

    for idx, file_path in enumerate(files, start=1):
        print(
            f"[{idx}/{len(files)}] Ingesting: {file_path.name}"
        )

        try:
            inserted = ingest_file(file_path)
            total_inserted += inserted or 0

        except Exception as e:
            print(
                f"  ERROR: Unexpected failure for "
                f"{file_path.name}: {e}"
            )
            traceback.print_exc()
            total_errors += 1

        print()

    print(f"{'=' * 60}")
    print("INGESTION COMPLETE")
    print(f"{'=' * 60}")
    print(f"Documents processed: {len(files)}")
    print(f"Total chunks inserted: {total_inserted}")

    if total_errors:
        print(f"Documents with errors: {total_errors}")

    print(f"{'=' * 60}\n")


if __name__ == "__main__":
    if len(sys.argv) > 1:
        target = Path(sys.argv[1])

        if target.is_dir():
            ingest_folder(str(target))

        elif target.is_file():
            print(f"Ingesting single file: {target.name}")
            result = ingest_file(target)
            print(f"Done. Inserted {result} chunks.")

        else:
            print(f"ERROR: Path does not exist: {target}")
            sys.exit(1)

    else:
        docs_folder = Path(__file__).parent / "documents"

        if docs_folder.exists():
            ingest_folder(str(docs_folder))

        else:
            print(
                f"ERROR: Default documents folder not found: "
                f"{docs_folder}"
            )
            sys.exit(1)